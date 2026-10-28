#!/usr/bin/env python3
"""domain-report.py - business-readable result of comparing ONE branch with a common set (gold baseline).

Turns the output of `gold.mjs compare` into a workbook in the same format as the L1 / L2 extraction workbook:
  Summary · Needs review · Missing in <X> · Matching · Extra in <X> · (Waived) · Common set

Every detail tab starts with API · Layer · Field · Check · Values · Runs · Validation, one colour band per API,
a thicker line where a new API starts, roomy rows and Runs coloured (Always green, If present blue, Conditional amber).

Usage (from automation-validation-extract/ - normally run by `npm run compare`):
  python3 domain-report.py --in out/.../report/validations.json --gold out/common/<name>/1_common/<name>_common.json \
      --compare out/.../compare/<branch>_vs_<name>_common.json.xlsx --branch draft-FIS12SL-2.3.0 --label FIS12SL-2.3.0 \
      --out out/.../<label>_vs_<name>.xlsx

Every number comes from the compare output; same input -> same output.
"""
import argparse, csv, json, collections, os, re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ap = argparse.ArgumentParser()
ap.add_argument('--in', dest='inp', required=True, help='validations.json from extract-validations.mjs')
ap.add_argument('--gold', required=True, help='common set: JSON from `gold.mjs partition` or CSV from build-gold')
ap.add_argument('--compare', required=True, help='<branch>_vs_<gold>.xlsx from gold.mjs compare')
ap.add_argument('--branch', required=True, help='branch that was compared, e.g. draft-FIS12SL-2.3.0')
ap.add_argument('--label', default=None, help='short name used in tab titles, e.g. FIS12SL-2.3.0')
ap.add_argument('--out', required=True, help='output .xlsx')
a = ap.parse_args()
X = a.label or a.branch

# ── inputs ──
if a.gold.endswith('.json'):   # { "L1": [...], "L2": [...] }
    _g = json.load(open(a.gold, encoding='utf-8'))
    gold = [{k: '' if v is None else str(v) for k, v in r.items()} for r in (_g if isinstance(_g, list) else _g.get('L1', []) + _g.get('L2', []))]
else:
    gold = list(csv.DictReader(open(a.gold, encoding='utf-8')))
gold_name = os.path.basename(a.gold).rsplit('.', 1)[0]
built_from = []   # the branches the common set was built from (summary.json of `npm run common`, or gold/index.yaml)
summ = [f for f in os.listdir(os.path.dirname(os.path.dirname(os.path.abspath(a.gold)))) if f.endswith('_summary.json')] if os.path.isdir(os.path.dirname(os.path.dirname(os.path.abspath(a.gold)))) else []
if summ:
    try: built_from = json.load(open(os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(a.gold))), summ[0]), encoding='utf-8')).get('use_cases', [])
    except Exception: pass
idx = os.path.join(os.path.dirname(os.path.abspath(a.gold)), 'index.yaml')
if not built_from and os.path.exists(idx):
    txt = open(idx, encoding='utf-8').read()
    for block in re.split(r'\n  - name:', txt):
        if re.search(r'file:\s*' + re.escape(os.path.basename(a.gold)) + r'\s*$', block, re.M):
            m = re.search(r'built_from:\s*\n((?:\s+- .+\n?)+)', block)
            if m: built_from = re.findall(r'- (\S+)', m.group(1))
data = json.load(open(a.inp, encoding='utf-8'))
sentence = {}   # plain-English sentence for every key, from the extractor output
for r in data['l1']: sentence.setdefault(('L1', r['branch'], r['api'], r['field'], r['ruleType']), r['validation'])
for r in data['l2']: sentence.setdefault(('L2', r['branch'], r['api'], r['field'], r['checkType']), r['validation'])
gold_sentence = {g['Key']: g.get('Validation', '') for g in gold}
l2_apis = {r['api'] for r in data['l2'] if r['branch'] == a.branch}   # APIs where the branch has real L2 checks

wbc = openpyxl.load_workbook(a.compare)
ws = wbc['Comparison']; hdr = [c.value for c in ws[1]]
rows = [dict(zip(hdr, [c.value if c.value is not None else '' for c in r])) for r in ws.iter_rows(min_row=2)]
for r in rows: r.setdefault('Key', ' | '.join([r['Layer'], r['API'], r['Field'], r['Check']]))   # same key as the common set
ORDER = ['search', 'on_search', 'select', 'on_select', 'init', 'on_init', 'confirm', 'on_confirm', 'status', 'on_status',
         'update', 'on_update', 'cancel', 'on_cancel', 'track', 'on_track', 'dynamic_form', 'issue', 'on_issue', 'issue_status', 'on_issue_status']
rank = lambda api: ORDER.index(api) if api in ORDER else 99
rows.sort(key=lambda r: (rank(r['API']), r['API'], r['Layer'], r['Field'], r['Check']))
text = lambda r: gold_sentence.get(r['Key']) or sentence.get((r['Layer'], a.branch, r['API'], r['Field'], r['Check']), '')

match = [r for r in rows if r['Status'] == 'OK']
miss = [r for r in rows if r['Status'] == 'MISSING']
waived = [r for r in rows if r['Status'] == 'WAIVED']
review = [r for r in rows if r['Status'] in ('BLOCKER', 'CONDITIONAL', 'DIFFERENT VALUES')]
extra = [r for r in rows if r['Status'] == 'EXTRA']

# ── plain wording ──
RESULT = {'BLOCKER': 'Blocker - skipped in more cases', 'CONDITIONAL': 'Runs in fewer cases', 'DIFFERENT VALUES': 'Different values'}
ACTION = {'BLOCKER': f'Must fix: remove the extra skip condition in {X}, or add the condition to the common set, or waive it with a reason',
          'CONDITIONAL': f'Check that the extra condition in {X} is intended; otherwise align it with the common set',
          'DIFFERENT VALUES': f'Align the allowed values of {X} with the common set, or waive the difference with a reason'}
def why_missing(r):
    if r['Layer'] == 'L2' and r['API'] not in l2_apis:
        return f'{X} has no L2 checks for {r["API"]} (no such step, or its validate code is an empty pass-through)'
    if r['Layer'] == 'L2': return f'{X} has {r["API"]} checks, but not this one'
    return f'{X} has no rule for this field'
def what_differs(r):
    if r['Status'] == 'DIFFERENT VALUES':
        g = set(str(r['Gold values']).split(', ')); b = set(str(r['Branch values']).split(', '))
        return f"{X} does not accept: {', '.join(sorted(g - b)) or '-'}" + (f"; {X} also accepts: {', '.join(sorted(b - g))}" if b - g else '')
    return str(r['What to do'])
def runs_short(s): return str(s or '').split(' (')[0]

# ── styling (same as the extraction workbook) ──
NAVY, FONT = 'FF1F3A5F', 'Arial'
GRID = Side(style='thin', color='FFD9DEE5'); SEP = Side(style='medium', color='FF8A97A8')
BANDS = ['FFFFFFFF', 'FFF3F6FA']
RUNS = {'Always': 'FF1E6B34', 'If present': 'FF1F5FA8', 'Conditional': 'FFB54708'}
TAB = {'Matching': 'FF70AD47', 'Missing': 'FFC00000', 'Review': 'FFFFC000', 'Extra': 'FF5B9BD5', 'Waived': 'FFA6A6A6', 'Common': 'FF1F3A5F'}
RESULT_FONT = {'Blocker - skipped in more cases': 'FFB42318', 'Runs in fewer cases': 'FFB54708', 'Different values': 'FFB54708'}
W = {'API': 16, 'Layer': 8, 'Field': 44, 'Check': 16, 'Values': 34, 'Runs': 36, 'Validation': 60, 'Result': 22,
     "What's different": 60, 'What to do': 45, f'Runs in {X}': 36, f'Values in {X}': 30, f'Where in {X}': 45, 'Code': 50,
     'Why missing': 46, 'Example (from another use case)': 60, 'Note': 34, 'Runs in common set': 36, f'Runs in {X} ': 36}
wb = openpyxl.Workbook()

def detail_sheet(name, cols, data, tab):
    s = wb.create_sheet(name[:31]); s.sheet_properties.tabColor = tab; s.sheet_view.showGridLines = False
    s.append(cols)
    for c in s[1]:
        c.font = Font(bold=True, color='FFFFFFFF', name=FONT, size=10); c.fill = PatternFill('solid', fgColor=NAVY)
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='left', indent=1)
    s.row_dimensions[1].height = 34
    for i, c in enumerate(cols, 1): s.column_dimensions[get_column_letter(i)].width = W.get(c, 18)
    band, prev = 0, None
    for d in data:
        s.append([d.get(c, '') for c in cols]); r = s.max_row
        new = d.get('API') != prev
        if new and prev is not None: band = 1 - band
        prev = d.get('API')
        lines = 1
        for i, c in enumerate(cols, 1):
            per = max(8, int((W.get(c, 18) - 2) * 1.15))
            lines = max(lines, sum(max(1, -(-len(p) // per)) for p in str(d.get(c, '') or '').split('\n')))
        s.row_dimensions[r].height = max(24, lines * 14 + 10)
        for i, c in enumerate(cols, 1):
            cell = s.cell(r, i)
            cell.font = Font(name=FONT, size=10, bold=(i == 1 and new), color='FF1F2933')
            cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='left', indent=1)
            cell.fill = PatternFill('solid', fgColor=BANDS[band])
            cell.border = Border(top=SEP if new and r > 2 else GRID, bottom=GRID, left=GRID, right=GRID)
            if c.startswith('Runs'):
                k = runs_short(cell.value)
                if k in RUNS: cell.font = Font(name=FONT, size=10, color=RUNS[k], bold=k != 'Always')
            if c == 'Result' and cell.value in RESULT_FONT: cell.font = Font(name=FONT, size=10, bold=True, color=RESULT_FONT[cell.value])
    s.freeze_panes = 'A2'
    if data: s.auto_filter.ref = f'A1:{get_column_letter(len(cols))}{len(data) + 1}'
    return s

FIRST = ['API', 'Layer', 'Field', 'Check', 'Values', 'Runs', 'Validation']
detail_sheet('Needs review', FIRST + ['Result', "What's different", 'What to do', f'Runs in {X}', f'Values in {X}'],
    [dict(r, Values=r['Gold values'], Runs=r['Gold runs'], Validation=text(r), Result=RESULT[r['Status']], **{"What's different": what_differs(r),
          'What to do': ACTION[r['Status']], f'Runs in {X}': r['Branch runs'], f'Values in {X}': r['Branch values']})
     for r in sorted(review, key=lambda r: (['BLOCKER', 'CONDITIONAL', 'DIFFERENT VALUES'].index(r['Status']), rank(r['API']), r['API'], r['Field']))], TAB['Review'])
detail_sheet(f'Missing in {X}', FIRST + ['Why missing', 'What to do'],
    [dict(r, Values=r['Gold values'], Runs=r['Gold runs'], Validation=text(r), **{'Why missing': why_missing(r),
          'What to do': f'Add it to the {X} {r["API"]} step(s), or waive it with a reason'}) for r in miss], TAB['Missing'])
detail_sheet('Matching', FIRST + [f'Runs in {X} '],
    [dict(r, Values=r['Gold values'], Runs=r['Gold runs'], Validation=text(r), **{f'Runs in {X} ': r['Branch runs']}) for r in match], TAB['Matching'])
detail_sheet(f'Extra in {X}', FIRST + ['Note'],
    [dict(r, Values=r['Branch values'], Runs=r['Branch runs'], Validation=text(r), Note=f'only in {X} - allowed (not in the common set)') for r in extra], TAB['Extra'])
if waived:
    detail_sheet('Waived', FIRST + ['What to do'], [dict(r, Values=r['Gold values'], Runs=r['Gold runs'], Validation=text(r), **{'What to do': r['What to do']}) for r in waived], TAB['Waived'])
detail_sheet(f'Common set ({len(gold)})', FIRST, sorted(gold, key=lambda g: (rank(g['API']), g['API'], g['Layer'], g['Field'])), TAB['Common'])

# ── Summary ──
s = wb.active; s.title = 'Summary'; s.sheet_view.showGridLines = False; s.sheet_properties.tabColor = NAVY
for col, w in zip('ABCDEFGH', [30, 12, 12, 12, 64, 22, 12, 12]): s.column_dimensions[col].width = w
n = lambda lst, layer=None: len([r for r in lst if layer is None or r['Layer'] == layer])
gl1 = sum(1 for g in gold if g['Layer'] == 'L1'); gl2 = len(gold) - gl1
cnt = collections.Counter(r['Status'] for r in review)
total = len(match) + len(miss) + len(review) + len(waived)
result = 'FAIL' if miss or cnt['BLOCKER'] else 'WARN' if review else 'PASS'
fill = lambda c, col: setattr(c, 'fill', PatternFill('solid', fgColor=col))
def put(r, vals, bold=False, size=10, color='FF1F2933', head=False, centre_from=2):
    for i, v in enumerate(vals, 1):
        c = s.cell(r, i, v)
        c.font = Font(name=FONT, size=size, bold=bold or head, color='FFFFFFFF' if head else color)
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center' if (i >= centre_from and isinstance(v, (int, float))) or (head and i > 1) else 'left', indent=0 if isinstance(v, (int, float)) else 1)
        if head: fill(c, NAVY)
        if v != '' or head: c.border = Border(top=GRID, bottom=GRID, left=GRID, right=GRID)
    s.row_dimensions[r].height = 34 if head else max(22, 15 * max(1, max(len(str(v)) // 70 + 1 for v in vals)) + 6)

put(1, [f'{X} vs common validations'], bold=True, size=16, color=NAVY); s.row_dimensions[1].height = 30
put(2, [f'Common set: {gold_name}  ·  {len(gold)} validations (L1 {gl1}, L2 {gl2})' + (f'  ·  built from {", ".join(built_from)}' if built_from else '')], color='FF667085')
put(3, [f'Branch compared: {a.branch}'], color='FF667085')
rc = s.cell(4, 1, f'Result: {result}'); rc.font = Font(name=FONT, size=13, bold=True, color={'FAIL': 'FFB42318', 'WARN': 'FFB54708', 'PASS': 'FF1E6B34'}[result])
s.row_dimensions[4].height = 26
r = 6
s.cell(r, 1, '1. Result').font = Font(name=FONT, size=12, bold=True, color=NAVY); r += 1
put(r, ['Result', 'L1', 'L2', 'Total', 'What it means', 'Tab'], head=True); r += 1
for label, lst, meaning, tab, key in [
        ('Matching', match, f'the common validation is in {X} and runs the same way', 'Matching', 'Matching'),
        ('Needs review', review, f'in {X}, but runs in fewer cases or with other values (blockers fail the build)', 'Needs review', 'Review'),
        (f'Missing in {X}', miss, f'the common validation is not in {X} - fails the build', f'Missing in {X}', 'Missing'),
        ('Waived', waived, 'accepted exceptions from the waivers file', 'Waived' if waived else '-', 'Waived'),
        (f'Extra in {X}', extra, f'only in {X} (not in the common set) - allowed', f'Extra in {X}', 'Extra')]:
    put(r, [label, n(lst, 'L1'), n(lst, 'L2'), len(lst), meaning, tab]); fill(s.cell(r, 1), {'Matching': 'FFE2F0D9', 'Review': 'FFFFF2CC', 'Missing': 'FFF8CBAD', 'Waived': 'FFEDEDED', 'Extra': 'FFDDEBF7'}[key]); r += 1
put(r, ['Common set total', gl1, gl2, len(gold), 'Matching + Needs review + Missing + Waived', ''], bold=True)
for i in range(1, 7): fill(s.cell(r, i), 'FFDCE6F2')
r += 2
s.cell(r, 1, '2. Needs review - details').font = Font(name=FONT, size=12, bold=True, color=NAVY); r += 1
put(r, ['Kind', 'L1', 'L2', 'Total', 'What it means', ''], head=True); r += 1
for k in ['BLOCKER', 'CONDITIONAL', 'DIFFERENT VALUES']:
    lst = [x for x in review if x['Status'] == k]
    put(r, [RESULT[k], n(lst, 'L1'), n(lst, 'L2'), len(lst), ACTION[k], '']); r += 1
r += 1
s.cell(r, 1, '3. Key findings').font = Font(name=FONT, size=12, bold=True, color=NAVY); r += 1
stub_missing = [x for x in miss if x['Layer'] == 'L2' and x['API'] not in l2_apis]
findings = [
    f"Field rules (L1): {n(match, 'L1')} of {gl1} common rules match; {n(miss, 'L1')} missing; {n(review, 'L1')} need review.",
    f"Flow code (L2): {n(match, 'L2')} of {gl2} common checks match; {n(miss, 'L2')} missing"
    + (f" ({len(stub_missing)} in APIs where {X} has no checks at all: {', '.join(sorted({x['API'] for x in stub_missing}))})" if stub_missing else '') + '.',
    f"{cnt['BLOCKER']} blocker(s): common checks {X} skips under an extra condition the common set does not have."
    + (" APIs: " + ', '.join(sorted({x['API'] for x in review if x['Status'] == 'BLOCKER'}, key=rank)) if cnt['BLOCKER'] else ''),
    f"{cnt['CONDITIONAL']} check(s) run in fewer cases than the common set; {cnt['DIFFERENT VALUES']} accept different values.",
    f"{len(extra)} checks only in {X} ({n(extra, 'L1')} L1, {n(extra, 'L2')} L2) - allowed.",
]
if a.label and any(a.label in b for b in built_from):
    findings.append(f'{X} is one of the use cases this common set was built from, so it is partly compared with its own validations.')
for i, t in enumerate(findings, 1):
    s.cell(r, 1, i).alignment = Alignment(horizontal='center', vertical='top'); s.cell(r, 1).font = Font(name=FONT, size=10, bold=True)
    c = s.cell(r, 2, t); c.font = Font(name=FONT, size=10); c.alignment = Alignment(wrap_text=True, vertical='top')
    s.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6); s.row_dimensions[r].height = 15 * (len(t) // 110 + 1) + 6; r += 1
for rr in (1, 2, 3, 4): s.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=6)
# Per API tab: one row per call with every result as its own column, Total row
p = wb.create_sheet('Per API'); p.sheet_view.showGridLines = False; p.sheet_properties.tabColor = NAVY
pc = ['API (call)', 'In common set', 'Matching', 'Needs review', 'of which blockers', f'Missing in {X}', f'Extra in {X}']
p.append(pc)
apis = sorted({x['API'] for x in rows}, key=lambda x: (rank(x), x))
cnt_api = lambda lst, api: len([x for x in lst if x['API'] == api])
prow = [[api, len([g for g in gold if g['API'] == api]), cnt_api(match, api), cnt_api(review, api), cnt_api([x for x in review if x['Status'] == 'BLOCKER'], api), cnt_api(miss, api), cnt_api(extra, api)] for api in apis]
prow.append(['Total', len(gold), len(match), len(review), cnt['BLOCKER'], len(miss), len(extra)])
for v in prow: p.append(v)
for i, w in enumerate([24, 14, 12, 14, 16, 18, 18], 1): p.column_dimensions[get_column_letter(i)].width = w
for ri, row in enumerate(p.iter_rows(), 1):
    p.row_dimensions[ri].height = 34 if ri == 1 else 22
    for ci, c in enumerate(row, 1):
        c.border = Border(top=GRID, bottom=GRID, left=GRID, right=GRID)
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='left' if ci == 1 and ri > 1 else 'center', indent=1 if ci == 1 else 0)
        if ri == 1: c.font = Font(name=FONT, size=10, bold=True, color='FFFFFFFF'); fill(c, NAVY)
        elif ri == p.max_row: c.font = Font(name=FONT, size=10, bold=True); fill(c, 'FFDCE6F2')
        else:
            c.font = Font(name=FONT, size=10, color='FFB42318' if (ci in (5, 6) and c.value) else 'FF1F2933', bold=bool(ci in (5, 6) and c.value))
            if ri % 2: fill(c, BANDS[1])
p.freeze_panes = 'A2'
wb.move_sheet('Per API', offset=-wb.index(wb['Per API']))
wb.move_sheet('Summary', offset=-wb.index(wb['Summary']))
wb.save(a.out)

assert total == len(gold), f'check failed: {total} rows vs {len(gold)} gold validations'
print(f'{a.branch}: matching={len(match)} missing={len(miss)} review={len(review)} '
      f'(blocker={cnt["BLOCKER"]}) waived={len(waived)} extra={len(extra)} | gold={len(gold)} ✓ -> {a.out}')
